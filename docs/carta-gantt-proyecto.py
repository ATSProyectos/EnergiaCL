from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

ARIAL = "Arial"
VERDE = "D9EAD3"      # relleno de barra
GRIS  = "808080"
AZUL  = "0000FF"
AMAR  = "FFFF00"

wb = Workbook()
ws = wb.active
ws.title = "Carta Gantt"

thin = Side(style="thin", color=GRIS)
med  = Side(style="medium", color="404040")
box  = Border(left=thin, right=thin, top=thin, bottom=thin)

ANIOS = [2027, 2028, 2029, 2030]
MESES = ["Ene","Feb","Mar","Abr","May","Jun","Jul","Ago","Sep","Oct","Nov","Dic"]

COL_MES0 = 5                      # primera columna de mes (E)
N_MESES  = 12 * len(ANIOS)
COL_PUNTOS = COL_MES0 + N_MESES   # columna "..."  (2031-2066)
COL_FIN    = COL_PUNTOS + 1       # columna "2067"

FILA_ANIO, FILA_MES, FILA_FECHA = 1, 2, 3
FILA0 = 4                         # primera fila de actividad

# ---------- Encabezados fijos ----------
for col, titulo, ancho in ((1,"Actividades",26), (2,"Inicio",12), (3,"Fin",12), (4,"Meses",8)):
    ws.merge_cells(start_row=FILA_ANIO, start_column=col, end_row=FILA_MES, end_column=col)
    c = ws.cell(row=FILA_ANIO, column=col, value=titulo)
    c.font = Font(name=ARIAL, size=11, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box
    ws.column_dimensions[get_column_letter(col)].width = ancho

# ---------- Encabezados de años / meses ----------
for i, anio in enumerate(ANIOS):
    ini = COL_MES0 + i * 12
    ws.merge_cells(start_row=FILA_ANIO, start_column=ini, end_row=FILA_ANIO, end_column=ini + 11)
    c = ws.cell(row=FILA_ANIO, column=ini, value=anio)
    c.font = Font(name=ARIAL, size=12, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.number_format = "0"
    for j, mes in enumerate(MESES):
        col = ini + j
        m = ws.cell(row=FILA_MES, column=col, value=mes)
        m.font = Font(name=ARIAL, size=7)
        m.alignment = Alignment(horizontal="center", vertical="center")
        m.border = box
        ws.cell(row=FILA_FECHA, column=col, value=date(anio, j + 1, 1)).number_format = "mmm-yyyy"
        ws.column_dimensions[get_column_letter(col)].width = 3.4
        ws.cell(row=FILA_ANIO, column=col).border = box

# ---------- Columnas "..." y 2067 ----------
for col, titulo, fecha, ancho in ((COL_PUNTOS, "...", date(2031, 1, 1), 4.5),
                                  (COL_FIN,   2067,  date(2067, 1, 1), 8)):
    ws.merge_cells(start_row=FILA_ANIO, start_column=col, end_row=FILA_MES, end_column=col)
    c = ws.cell(row=FILA_ANIO, column=col, value=titulo)
    c.font = Font(name=ARIAL, size=12, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box
    ws.cell(row=FILA_MES, column=col).border = box
    ws.cell(row=FILA_FECHA, column=col, value=fecha).number_format = "mmm-yyyy"
    ws.column_dimensions[get_column_letter(col)].width = ancho

# ---------- Actividades (celdas de entrada: Inicio / Fin) ----------
ACTIVIDADES = [
    ("Fase de Estudio",      date(2027, 1, 1),  date(2029, 10, 1)),
    ("Fase de Construcción", date(2029, 11, 1), date(2030, 10, 1)),
    ("Fase Operación",       date(2030, 11, 1), date(2066, 12, 1)),
    ("Fase Cierre",          date(2067, 1, 1),  date(2067, 12, 1)),
]

for k, (nombre, ini, fin) in enumerate(ACTIVIDADES):
    f = FILA0 + k
    ws.row_dimensions[f].height = 20
    a = ws.cell(row=f, column=1, value=nombre)
    a.font = Font(name=ARIAL, size=11)
    a.alignment = Alignment(horizontal="center", vertical="center")
    a.border = box
    for col, val in ((2, ini), (3, fin)):
        c = ws.cell(row=f, column=col, value=val)
        c.font = Font(name=ARIAL, size=10, color=AZUL)
        c.fill = PatternFill("solid", fgColor=AMAR)
        c.number_format = "mmm-yyyy"
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = box
    d = ws.cell(row=f, column=4, value=f"=IF(OR($B{f}=\"\",$C{f}=\"\"),\"\",DATEDIF($B{f},$C{f},\"m\")+1)")
    d.font = Font(name=ARIAL, size=10)
    d.alignment = Alignment(horizontal="center", vertical="center")
    d.border = box
    for col in range(COL_MES0, COL_FIN + 1):
        ws.cell(row=f, column=col).border = box

# ---------- Formato condicional: la barra se dibuja sola ----------
ULTIMA = FILA0 + len(ACTIVIDADES) - 1
rango = f"{get_column_letter(COL_MES0)}{FILA0}:{get_column_letter(COL_FIN)}{ULTIMA}"
regla = FormulaRule(
    formula=[f'AND($B{FILA0}<>"",$C{FILA0}<>"",{get_column_letter(COL_MES0)}${FILA_FECHA}>=$B{FILA0},'
             f'{get_column_letter(COL_MES0)}${FILA_FECHA}<=$C{FILA0})'],
    fill=PatternFill(start_color=VERDE, end_color=VERDE, fill_type="solid"),
    stopIfTrue=False,
)
ws.conditional_formatting.add(rango, regla)

# ---------- Leyenda ----------
leyenda = [
    "Cómo editar esta carta Gantt:",
    "1. Escriba las fechas en las columnas Inicio y Fin (celdas amarillas, texto azul). La barra verde se dibuja automáticamente.",
    "2. Formato de fecha: use el primer día del mes (ej. 01-11-2029). Se muestra como mmm-aaaa.",
    "3. Para agregar una actividad: copie una fila existente hacia abajo (así se copia también el formato condicional) y cambie nombre y fechas.",
    "4. La columna 'Meses' se calcula sola con DATEDIF; no la escriba a mano.",
    "5. La columna '...' representa los años 2031-2066 (operación continua) y la columna '2067' el año de cierre.",
    "",
    "Fuente de los datos: cronograma entregado por el usuario (imagen de la carta Gantt del proyecto). Las fechas son supuestos del proyecto, no hay otra fuente externa.",
]
for i, txt in enumerate(leyenda):
    c = ws.cell(row=ULTIMA + 2 + i, column=1, value=txt)
    c.font = Font(name=ARIAL, size=9, bold=(i == 0))
    c.alignment = Alignment(horizontal="left", vertical="center")

ws.row_dimensions[FILA_FECHA].hidden = True   # fila auxiliar de fechas
ws.row_dimensions[FILA_ANIO].height = 20
ws.row_dimensions[FILA_MES].height = 16
ws.freeze_panes = "E4"
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.save("/home/user/EnergiaCL/docs/carta-gantt-proyecto.xlsx")
print("ok")
